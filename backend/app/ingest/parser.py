"""Парсинг документов в чистый текст (ARCHITECTURE.md §4, шаг 1).

Форматы кейса: PDF, DOCX, PPTX (русскоязычные), плюс .md/.txt. На выходе —
плоский текст + `content_hash = sha256(text)` для идемпотентного пропуска
неизменённых документов (инвариант №4). `full_text` в узле Document НЕ хранится
(§3.2) — writer кладёт его файлом на диск; здесь мы только извлекаем текст.

Числа/сущности отсюда НЕ извлекаются — этим занимается extractor (§4.1) поверх
чанков. Задача парсера — дать один связный текст на документ, максимально близкий
к читаемому: постранично для PDF, по параграфам/таблицам для DOCX, по слайдам для
PPTX. Нормализация лёгкая (NBSP→пробел, схлопывание пустых строк) — агрессивно
чистить нельзя, иначе validator (§4.2) не найдёт `quote` подстрокой.
"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from pathlib import Path


class UnsupportedFormat(Exception):
    """Расширение файла не входит в поддерживаемые (.pdf/.docx/.pptx/.md/.txt)."""


@dataclass
class ParsedDoc:
    """Результат парсинга одного файла (§4, шаг 1)."""

    text: str
    content_hash: str
    source_path: str
    suffix: str


# Неразрывный пробел (NBSP, U+00A0) и узкий неразрывный (U+202F) → обычный пробел:
# из PDF/DOCX они лезут массово и ломают regex-сверку чисел в validator (§4.2).
_NBSP_CHARS = "  "
_NBSP_RE = re.compile(f"[{_NBSP_CHARS}]")
# Три и более подряд идущих переводов строки (с любыми пробелами между ними)
# схлопываем максимум до двух — граница абзаца для chunker (§4, шаг 3).
_BLANK_LINES_RE = re.compile(r"\n[ \t]*\n[ \t]*(?:\n[ \t]*)+")


def normalize_text(text: str) -> str:
    """Лёгкая нормализация (§4, шаг 1): NBSP→пробел, схлопывание >2 пустых строк.

    Намеренно консервативна: сохраняет исходные подстроки чисел и цитат, чтобы
    validator нашёл `quote`/`value_raw` в тексте чанка (§4.2, инвариант №1).
    """
    text = _NBSP_RE.sub(" ", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # хвостовые пробелы на строках убираем — они мешают детекции границ абзацев
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    text = _BLANK_LINES_RE.sub("\n\n", text)
    return text.strip()


def _hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


# ---------------------------------------------------------------------------
# Парсеры по форматам
# ---------------------------------------------------------------------------

def processed_output_path(src: Path) -> Path | None:
    """Путь .md-зеркала распарсенного текста (04.07): corpus/<...>/x.docx →
    processed_corpus/<...>/x.md, с сохранением вложенности папок.

    Ищем предка с именем «corpus» (CLI может указывать на подкаталог —
    относительный путь всё равно строим от корня corpus). Файл вне corpus/ → None.
    """
    for parent in src.resolve().parents:
        if parent.name == "corpus":
            rel = src.resolve().relative_to(parent)
            return parent.parent / "processed_corpus" / rel.with_suffix(".md")
    return None


def write_processed_md(parsed: "ParsedDoc", src: Path) -> Path | None:
    """Сохранить распарсенный текст .md-зеркалом в processed_corpus/ (04.07).

    Детерминированный выход (front-matter без дат): одинаковый вход → одинаковый
    файл байт-в-байт, повторные прогоны не создают шума. Ошибка записи не должна
    валить импорт — вызывающий оборачивает в try/except.
    """
    out = processed_output_path(src)
    if out is None:
        return None
    out.parent.mkdir(parents=True, exist_ok=True)
    front = (
        "---\n"
        f"source: {src.resolve()}\n"
        f"content_hash: {parsed.content_hash}\n"
        "---\n\n"
    )
    out.write_text(front + parsed.text, encoding="utf-8")
    return out


def _pdf_table_ok(rows: list[list[str]]) -> bool:
    """Фильтр ложных детекций find_tables на презентационных PDF (03.07).

    Слайдовые раскладки детектор принимает за «таблицы» и порождает мусорные пары
    «предложение = предложение». Настоящая таблица данных: >= 3 строк, >= 2 колонок,
    >= 60% ячеек заполнено, ячейки КОРОТКИЕ (сред. <= 40 символов — не абзацы)
    и >= 30% ячеек с цифрами (данные, а не текстовая вёрстка).
    """
    if len(rows) < 3 or len(rows[0]) < 2:
        return False
    cells = [(c or "").strip() for r in rows for c in r]
    filled = [c for c in cells if c]
    if not filled or len(filled) / len(cells) < 0.6:
        return False
    avg_len = sum(len(c) for c in filled) / len(filled)
    digity = sum(1 for c in filled if any(ch.isdigit() for ch in c)) / len(filled)
    return avg_len <= 40 and digity >= 0.3


def _parse_pdf(path: Path) -> str:
    """PDF через PyMuPDF (fitz), постранично (§4, шаг 1).

    Таблицы: page.find_tables() → header-aware линеаризация (как DOCX/PPTX, 03.07).
    Сырой текст страницы уже содержит таблицу «кашей из колонок» — линеаризованная
    версия добавляется ПОСЛЕ текста страницы (дубль фактов схлопывает дедуп writer'а,
    зато цитаты становятся непрерывными). Ложные детекции на слайдах режет
    _pdf_table_ok; любые ошибки детектора не валят парсинг страницы.
    """
    import fitz  # PyMuPDF; импорт локальный — не тянуть зависимость в остальной код

    parts: list[str] = []
    with fitz.open(path) as doc:
        for page in doc:
            parts.append(page.get_text("text"))
            try:
                for table in page.find_tables():
                    rows = [
                        [(c if isinstance(c, str) else "") for c in row]
                        for row in table.extract()
                    ]
                    if _pdf_table_ok(rows):
                        lin = linearize_table(rows)
                        if lin.strip():
                            parts.append("[Таблица]\n" + lin)
            except Exception:  # noqa: BLE001 — детектор таблиц не должен валить импорт
                continue
    return "\n\n".join(parts)


def linearize_table(rows: list[list[str]]) -> str:
    """Header-aware линеаризация таблицы (03.07): каждая пара «колонка = значение»
    становится НЕПРЕРЫВНОЙ подстрокой текста.

    Построчная сериализация («| Cu+Ni | Pd … \\n Низкокремнистый | 3,14-7,24 | …»)
    разносила шапку и значение на десятки символов — модель извлекает по колонкам,
    цитата не сходится, validator ставил «quote не найдена» на каждый табличный факт.
    Теперь: «Низкокремнистый: Cu+Ni = 3,14-7,24; Pd = 16,0-26,1; …» — и цитируемо,
    и модели проще связывать параметр со значением.

    Правила: первая строка — шапка (если в таблице >= 2 строк); пустая ячейка шапки
    над первой колонкой = колонка меток строк; пары строятся позиционно (zip по
    короткому); таблица из одной строки — прежний « | »-формат.
    """
    clean = [[(c or "").strip() for c in r] for r in rows if any((c or "").strip() for c in r)]
    if not clean:
        return ""
    if len(clean) == 1:
        return " | ".join(c for c in clean[0] if c)

    header, data = clean[0], clean[1:]

    # Двухэтажная шапка (merged-ячейки: «Содержание в сплаве, % масс.» ×3 + Cu/Ni/Fe
    # во второй строке — python-docx размножает объединённые ячейки по колонкам).
    # Признаки: вторая строка нечисловая, в первой есть соседние дубли/пустоты.
    def _headerish(row: list[str]) -> bool:
        cells = [c for c in row if c]
        digity = sum(1 for c in cells if any(ch.isdigit() for ch in c))
        return bool(cells) and digity / len(cells) < 0.3

    def _merged_artifacts(row: list[str]) -> bool:
        return any(a and a == b for a, b in zip(row, row[1:])) or not all(row)

    if len(clean) >= 3 and len(clean[1]) == len(header) and _headerish(clean[1]) \
            and _merged_artifacts(header):
        header = [
            (f"{h1}, {h2}" if h2 and h2 != h1 and h1 else (h1 or h2))
            for h1, h2 in zip(header, clean[1])
        ]
        data = clean[2:]
    label_col = not header[0]  # пустая шапка первой колонки → это метки строк
    out: list[str] = []
    for row in data:
        if label_col:
            label, values = (row[0] or "строка"), row[1:]
            cols = header[1:]
        else:
            label, values = None, row
            cols = header
        pairs = [f"{h} = {v}" for h, v in zip(cols, values) if h and v and h != v]
        if not pairs:
            # шапка не спарилась (склеенные ячейки) — не теряем данные, старый формат
            out.append(" | ".join(c for c in row if c))
            continue
        out.append((f"{label}: " if label else "") + "; ".join(pairs) + ".")
    return "\n".join(out)


def _parse_docx(path: Path) -> str:
    """DOCX через python-docx: параграфы + таблицы (header-aware линеаризация).

    python-docx не отдаёт параграфы и таблицы в порядке их следования в документе,
    поэтому идём по XML-телу body и разбираем элементы (`w:p` / `w:tbl`) по порядку —
    иначе таблицы всплывают в конце и рвут контекст для extractor.
    """
    from docx import Document as _Docx
    from docx.document import Document as _DocDocument
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = _Docx(str(path))
    parts: list[str] = []

    def _table_text(table: Table) -> str:
        return linearize_table(
            [[cell.text for cell in row.cells] for row in table.rows]
        )

    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            para = Paragraph(child, doc)
            txt = para.text.strip()
            if txt:
                parts.append(txt)
        elif child.tag == qn("w:tbl"):
            table = Table(child, doc)
            txt = _table_text(table)
            if txt.strip():
                parts.append(txt)

    # Заголовки/колонтитулы python-docx в body не отдаёт — для кейса это не критично.
    _ = _DocDocument  # держим импорт осмысленным для читателя
    return "\n\n".join(parts)


def _parse_pptx(path: Path) -> str:
    """PPTX через python-pptx: текст всех shape по слайдам, префикс «Слайд N:» (§4)."""
    from pptx import Presentation

    prs = Presentation(str(path))
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        chunks: list[str] = [f"Слайд {i}:"]
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                # Таблицы на слайдах — отдельный контейнер (не text_frame).
                if getattr(shape, "has_table", False):
                    tbl = linearize_table(
                        [[cell.text for cell in row.cells] for row in shape.table.rows]
                    ).strip()
                    if tbl:
                        chunks.append(tbl)
                continue
            txt = shape.text_frame.text.strip()
            if txt:
                chunks.append(txt)
        if len(chunks) > 1:  # не добавляем пустой слайд (только префикс)
            parts.append("\n".join(chunks))
    return "\n\n".join(parts)


def _parse_text(path: Path) -> str:
    """.md/.txt как есть (§4, шаг 1). UTF-8 с игнором битых байтов."""
    return path.read_text(encoding="utf-8", errors="ignore")


# Диспетчер по расширению (нижний регистр). .docm — Word с макросами, тот же формат.
_PARSERS = {
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".docm": _parse_docx,
    ".pptx": _parse_pptx,
    ".md": _parse_text,
    ".txt": _parse_text,
}


def parse_file(path: Path) -> ParsedDoc:
    """Парсит файл в `ParsedDoc` (§4, шаг 1). `content_hash = sha256(нормализованный text)`.

    Расширение вне карты → UnsupportedFormat (fail fast, инвариант №9): молчаливого
    пропуска формата нет, интеграционный агент решает, что делать.
    """
    path = Path(path)
    suffix = path.suffix.lower()
    parser = _PARSERS.get(suffix)
    if parser is None:
        raise UnsupportedFormat(f"неподдерживаемый формат: {path.suffix} ({path.name})")

    raw = parser(path)
    text = normalize_text(raw)
    return ParsedDoc(
        text=text,
        content_hash=_hash(text),
        source_path=str(path),
        suffix=suffix,
    )
